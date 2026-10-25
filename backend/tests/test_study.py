import json

import pytest
from pptx import Presentation

from app import embeddings, llm
from app.slides import SlideError, extract_slides
from app.study import CARD_COUNTS, DocumentError, detect_kind, fit_source, parse_cards, retrieve_units


def _pptx(path, slides: list[tuple[str, str, str]]) -> bytes:
    deck = Presentation()
    layout = deck.slide_layouts[1]
    for title, body, notes in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    deck.save(str(path))
    return path.read_bytes()


def test_extract_slides_keeps_body_and_notes(tmp_path):
    path = tmp_path / "deck.pptx"
    path.write_bytes(_pptx(path, [("Mitosis", "Cells divide into two.", "Mention checkpoints.")]))
    count, slides = extract_slides(path)
    assert count == 1
    assert slides[0][0] == 1
    assert "Cells divide into two." in slides[0][1]
    assert "Speaker notes: Mention checkpoints." in slides[0][1]


def test_extract_slides_skips_blank_and_rejects_garbage(tmp_path):
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    titled = deck.slides.add_slide(deck.slide_layouts[1])
    titled.shapes.title.text = "Only this slide"
    deck.save(str(path))
    count, slides = extract_slides(path)
    assert count == 2
    assert [n for n, _ in slides] == [2]

    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"PK\x03\x04not a presentation")
    with pytest.raises(SlideError):
        extract_slides(bad)


def test_detect_kind():
    assert detect_kind("notes.pdf", b"%PDF-1.7") == "pdf"
    assert detect_kind("deck.pptx", b"PK\x03\x04") == "pptx"
    with pytest.raises(DocumentError, match="pptx"):
        detect_kind("old.ppt", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    with pytest.raises(DocumentError, match="PDF"):
        detect_kind("photo.png", b"\x89PNG")


def test_fit_source_keeps_every_label_under_the_limit():
    units = [(i, "word " * 500) for i in range(1, 6)]
    text = fit_source(units, "Slide", limit=800)
    assert len(text) <= 800
    for i in range(1, 6):
        assert f"[Slide {i}]" in text


def test_retrieve_units_prefers_query_matches_and_keeps_document_order():
    units = [(i, f"chunk {i}") for i in range(1, 7)]
    vectors = [
        [1.0, 0.0],
        [0.0, 1.0],
        [0.2, 0.2],
        [0.1, 0.0],
        [0.9, 0.1],
        [0.0, 0.2],
    ]
    picked = retrieve_units(units, vectors, [[1.0, 0.0], [0.0, 1.0]], k=4)
    numbers = [n for n, _ in picked]
    assert 5 in numbers and 2 in numbers
    assert numbers == sorted(numbers)
    assert len(picked) == 4


def test_card_counts():
    assert CARD_COUNTS == (3, 5, 8, 12, 15)


def test_parse_cards_accepts_fenced_json_and_drops_blanks():
    raw = """```json
    [
      {"front": " What is ATP? ", "back": "Cell energy currency.", "source": "Page 2"},
      {"front": "What is ATP?", "back": "duplicate"},
      {"question": "", "answer": "missing"},
      {"term": "Osmosis", "definition": "Water moves across a membrane.", "slide": "Slide 4"}
    ]
    ```"""
    cards = parse_cards(raw, limit=10)
    assert cards == [
        {"front": "What is ATP?", "back": "Cell energy currency.", "source": "Page 2"},
        {"front": "Osmosis", "back": "Water moves across a membrane.", "source": "Slide 4"},
    ]
    assert len(parse_cards(raw, limit=1)) == 1


def test_parse_cards_reads_object_wrapper():
    raw = json.dumps({"cards": [{"front": "Q", "back": "A", "source": "Page 1"}]})
    assert parse_cards(raw, 4)[0]["front"] == "Q"


def test_parse_cards_rejects_prose():
    with pytest.raises(ValueError):
        parse_cards("I could not make cards.", 8)


def test_study_pdf_returns_cards(client, pdf_bytes, monkeypatch):
    calls = {"passages": 0, "queries": 0}

    def passages(texts):
        calls["passages"] += 1
        assert texts and "mitochondria" in texts[0]
        return [[1.0, 0.0] for _ in texts]

    def query(text):
        calls["queries"] += 1
        assert text
        return [1.0, 0.0]

    async def fake(messages, transport=None, max_tokens=None, json_schema=None, **kwargs):
        assert max_tokens and json_schema
        assert "retrieved excerpts" in messages[1]["content"]
        assert "[Page 1]" in messages[1]["content"]
        assert "mitochondria" in messages[1]["content"]
        return json.dumps(
            [{"front": "Where is ATP made?", "back": "In mitochondria.", "source": "Page 1"}]
        )

    monkeypatch.setattr(embeddings, "embed_passages", passages)
    monkeypatch.setattr(embeddings, "embed_query", query)
    monkeypatch.setattr(llm, "complete", fake)
    data = pdf_bytes(["The mitochondria make ATP."])
    r = client.post(
        "/api/study",
        files={"file": ("bio.pdf", data, "application/pdf")},
        data={"count": "8"},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["kind"] == "pdf"
    assert body["units"] == 1
    assert body["cards"][0]["back"] == "In mitochondria."
    assert calls == {"passages": 1, "queries": 3}
    assert client.db["doc"] is None


def test_study_pptx_returns_cards(client, tmp_path, monkeypatch):
    async def fake(messages, transport=None, max_tokens=None, json_schema=None):
        assert "[Slide 1]" in messages[1]["content"]
        return json.dumps([{"front": "Topic", "back": "Cells divide.", "source": "Slide 1"}])

    monkeypatch.setattr(llm, "complete", fake)
    data = _pptx(tmp_path / "lecture.pptx", [("Mitosis", "Cells divide.", "")])
    r = client.post(
        "/api/study",
        files={
            "file": (
                "lecture.pptx",
                data,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
        data={"count": "8"},
    )
    assert r.status_code == 200
    assert r.json()["kind"] == "pptx"
    assert r.json()["cards"][0]["source"] == "Slide 1"


def test_study_rejects_counts_outside_the_picker(client, pdf_bytes):
    data = pdf_bytes(["Cells have membranes."])
    for count in ("2", "4", "16", "24"):
        r = client.post(
            "/api/study",
            files={"file": ("bio.pdf", data, "application/pdf")},
            data={"count": count},
        )
        assert r.status_code == 422


def test_study_rejects_legacy_ppt_and_empty_pdf(client, pdf_bytes):
    old = client.post("/api/study", files={"file": ("old.ppt", b"\xd0\xcf\x11\xe0rest", "application/vnd.ms-powerpoint")})
    assert old.status_code == 400
    empty = client.post("/api/study", files={"file": ("blank.pdf", pdf_bytes([""]), "application/pdf")})
    assert empty.status_code == 422


def test_study_reports_bad_model_output(client, pdf_bytes, monkeypatch):
    async def fake(messages, transport=None, max_tokens=None, json_schema=None):
        return "sorry, no json"

    monkeypatch.setattr(llm, "complete", fake)
    r = client.post("/api/study", files={"file": ("bio.pdf", pdf_bytes(["Cells have membranes."]), "application/pdf")})
    assert r.status_code == 502
